"""Unit tests for ingest.parse — DOCX paragraphs + tables to Markdown."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from src.ingest.parse import parse_docx


def _write_docx(path: Path, build) -> Path:
    doc = Document()
    build(doc)
    doc.save(path)
    return path


def test_parse_empty_docx(tmp_path: Path) -> None:
    path = _write_docx(tmp_path / "empty.docx", lambda doc: None)
    assert parse_docx(path) == ""


def test_parse_paragraphs_joined_by_newline(tmp_path: Path) -> None:
    def build(doc: Document) -> None:
        doc.add_paragraph("第一段景点介绍")
        doc.add_paragraph("第二段开放时间")

    path = _write_docx(tmp_path / "paras.docx", build)
    assert parse_docx(path) == "第一段景点介绍\n第二段开放时间"


def test_parse_table_as_markdown(tmp_path: Path) -> None:
    def build(doc: Document) -> None:
        doc.add_paragraph("票价说明")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "票种"
        table.rows[0].cells[1].text = "价格"
        table.rows[1].cells[0].text = "成人"
        table.rows[1].cells[1].text = "399"

    path = _write_docx(tmp_path / "table.docx", build)
    text = parse_docx(path)
    assert "票价说明" in text
    assert "| 票种 | 价格 |" in text
    assert "|---|---|" in text or "|---|---|" in text.replace(" ", "")
    assert "| 成人 | 399 |" in text
